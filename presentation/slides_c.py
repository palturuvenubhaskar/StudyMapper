"""Slides 13-18."""
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.ns import qn
from deck_lib import *


def slide_13(prs):
    s = base(prs, "11 · Evidence", "What is verified — and what is not",
             "Results are reported at the level they were measured: build health and functional completeness.",
             section="Evidence", page=13)
    kpis = [
        ("47/47", "automated tests passing", "14 test files", GREEN),
        ("6.21 s", "production build", "Vite production build", TEAL),
        ("40", "database tables", "Dexie / IndexedDB", BLUE),
        ("46", "routed views", "React single-page app", AMBER),
    ]
    for i, (num, label, sub, col) in enumerate(kpis):
        x = 0.62 + i * 3.06
        card(s, x, 1.92, 2.86, 1.14, fill=WHITE, line=LINE, accent=col, accent_w=0.22)
        add_text(s, x + 0.20, 2.04, 1.6, 0.42, num, size=22, color=INK, bold=True)
        add_text(s, x + 0.20, 2.48, 2.45, 0.2, label, size=8.6, color=BODY, bold=True)
        add_text(s, x + 0.20, 2.68, 2.45, 0.2, sub, size=7.6, color=MUTED)
    add_text(s, 0.62, 3.20, 12.1, 0.2,
             "SEEDED DEMO DATASET   ·   38 topics   ·   18 note sets   ·   8 questions   ·   10-day plan   ·   12-day streak",
             size=7.6, color=MUTED, bold=True)

    rows = [
        ("Area", "Evidence"),
        ("Automated suite", "47 tests across 14 files — all passing; 6.21 s production build"),
        ("Data layer", "40 Dexie tables created, seeded and queried"),
        ("AI service", "Six-model fallback tested; streaming and JSON repair implemented"),
        ("Interface", "19 screens captured from the running application"),
        ("Deployment", "Production build verified; live demo reachable"),
    ]
    table(s, 0.62, 3.50, 7.55, [1.85, 5.70], rows, row_h=0.44, head_h=0.36, size=9.0)

    rect(s, 8.34, 3.50, 4.38, 2.52, fill=WHITE, line=LINE, radius=0.04)
    cd = CategoryChartData()
    cd.categories = ["Passed", "Not passing"]
    cd.add_series("Tests", (47, 0))
    gf = s.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(8.50), Inches(3.62),
                            Inches(1.62), Inches(1.62), cd)
    ch = gf.chart
    ch.has_title = False
    ch.has_legend = False
    ch.font.size = Pt(8)
    plot = ch.plots[0]
    plot.has_data_labels = False
    for idx, col in enumerate((TEAL, LINE)):
        pt = plot.series[0].points[idx]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = col
        pt.format.line.color.rgb = WHITE
        pt.format.line.width = Pt(1)
    hs = ch._chartSpace.find(".//" + qn("c:holeSize"))
    if hs is not None:
        hs.set("val", "72")
    add_text(s, 8.74, 4.14, 0.96, 0.3, "100%", size=13, color=INK, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, 8.70, 4.44, 1.04, 0.2, "PASS RATE", size=6.4, color=MUTED, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, 10.32, 3.86, 2.2, 0.34, "47 / 47", size=17, color=INK, bold=True)
    add_text(s, 10.32, 4.22, 2.25, 0.2, "tests passing", size=8.2, color=BODY)
    add_text(s, 10.32, 4.46, 2.25, 0.2, "0 failures", size=8.2, color=GREEN, bold=True)

    add_text(s, 8.50, 5.38, 2.0, 0.2, "BUILD PIPELINE", size=6.4, color=MUTED, bold=True)
    arrow(s, 8.61, 5.64, 11.67, 5.64, color=LINE, width=1.4)
    steps = ["Install", "Test", "Build", "Preview"]
    for i, t in enumerate(steps):
        x = 8.50 + i * 1.02
        circle(s, x, 5.53, 0.22, fill=TEAL, text="✓", tcolor=WHITE, size=6)
        add_text(s, x - 0.12, 5.78, 0.9, 0.2, t, size=6.6, color=BODY, align=PP_ALIGN.CENTER)

    rect(s, 0.62, 6.24, 12.10, 0.58, fill=AMBER_SOFT, line=None, radius=0.08)
    add_text(s, 0.82, 6.33, 1.4, 0.2, "INTERPRETATION", size=6.8, color=AMBER, bold=True)
    add_text(s, 0.82, 6.48, 11.6, 0.3,
             "This evidences functional completeness and build health. It does not yet evidence learning gains — a classroom study is future work.",
             size=8.6, color=BODY)
    notes(s, "Be precise about what these numbers mean. The test count, build time, table count and route count are measured from this build. The table lists what each area evidences. The donut is only a pass-rate visual, not a research result. The most important sentence is the amber band: this project proves the system works and is complete against its own scope; it does not prove that students learn more. Say that openly — it strengthens the work, because it shows the team understands the difference between engineering validation and educational evaluation.")


def slide_14(prs):
    s = base(prs, "12 · Discussion", "What it contributes — and what it does not claim",
             "A credible project states its boundaries as clearly as its features.",
             section="Evidence", page=14)
    add_text(s, 0.62, 1.94, 6.0, 0.2, "WHAT IT CONTRIBUTES", size=7.8, color=TEAL, bold=True)
    contrib = [
        ("A closed loop, not a bundle",
         "Planning, content, practice and measurement share one data model, so evidence actually changes the plan."),
        ("AI with a safety net",
         "Six-model fallback, streaming and JSON repair; generated material stays editable, never opaque."),
        ("Local-first by design",
         "The core loop runs offline; identity is optional and never required for study."),
    ]
    for i, (t, d) in enumerate(contrib):
        y = 2.22 + i * 0.80
        rect(s, 0.62, y + 0.03, 0.05, 0.52, fill=TEAL)
        add_text(s, 0.82, y, 5.9, 0.22, t, size=10.6, color=INK, bold=True)
        add_text(s, 0.82, y + 0.24, 5.95, 0.46, d, size=8.8, color=MUTED, line=1.18)

    rect(s, 8.28, 1.94, 4.44, 1.98, fill=ROSE_SOFT, line=None, radius=0.05)
    add_text(s, 8.52, 2.06, 3.9, 0.2, "WHAT IT DOES NOT CLAIM", size=7.8, color=ROSE, bold=True)
    add_text(s, 8.52, 2.30, 3.98, 1.12,
             "·  no measured learning-outcome study\n·  no multi-institution deployment\n·  no teacher or cohort dashboard\n·  no native mobile application",
             size=8.6, color=BODY, line=1.30)
    add_text(s, 8.52, 3.48, 3.95, 0.22, "These are scope decisions, not missing pieces.",
             size=7.4, color=ROSE, italic=True)

    add_text(s, 0.62, 4.62, 6.0, 0.2, "DELIBERATE TRADE-OFFS", size=7.8, color=TEAL, bold=True)
    rows = [
        ("Decision", "Why it was taken", "What it costs"),
        ("Local-first storage", "Privacy, instant reads and offline study", "No cross-device sync yet"),
        ("AI-assisted authoring", "Cuts content preparation time", "Output needs review; free models vary"),
        ("Breadth over depth", "Covers the whole study loop", "Each module is intentionally shallow"),
    ]
    table(s, 0.62, 4.84, 12.10, [2.45, 4.85, 4.80], rows, row_h=0.44, head_h=0.34, size=9.0)
    add_text(s, 0.62, 6.56, 12.1, 0.24,
             "The goal was a coherent system a student could use end-to-end, not a deep implementation of any single module.",
             size=8.6, color=MUTED, italic=True)
    notes(s, "This slide is for the evaluators. The contributions are stated as mechanisms, not adjectives: a shared data model, an AI layer that degrades gracefully, and local-first storage. The rose panel is deliberately prominent — no learning-outcome study, no institutional deployment, no teacher dashboard, no native app. Then the trade-off table explains why: local-first buys privacy and offline study at the cost of sync; AI authoring buys speed at the cost of review; breadth buys a complete loop at the cost of depth. If a reviewer asks what you would deepen first, point to the planner and the weakness model.")


def slide_15(prs):
    s = base(prs, "13 · Conclusion", "A complete, working learning loop",
             "Not a prototype of one — every part of the loop shown in this deck runs in the deployed build.",
             section="Close", page=15)
    pillars = [
        ("ADAPTIVE", TEAL, "Plans recompute from the exam date, topic difficulty and measured weakness.",
         "Regenerate on demand from the latest scores and remaining days."),
        ("EXPLAINABLE", BLUE, "Every score and schedule item traces back to stored attempts — no black-box advice.",
         "A topic is flagged weak at 40+; the window is the last seven days."),
        ("LOCAL-FIRST", AMBER, "The core loop runs offline; AI and sign-in are optional services.",
         "Notes, plans and attempts live in IndexedDB, not on a server."),
    ]
    for i, (name, col, d, detail) in enumerate(pillars):
        y = 2.02 + i * 1.34
        rect(s, 0.62, y, 0.05, 0.94, fill=col)
        add_text(s, 0.84, y - 0.02, 5.9, 0.2, name, size=8.2, color=col, bold=True)
        add_text(s, 0.84, y + 0.19, 6.15, 0.46, d, size=10.6, color=INK, bold=True, line=1.14)
        add_text(s, 0.84, y + 0.64, 6.12, 0.3, detail, size=8.4, color=MUTED, italic=True, line=1.12)

    rect(s, 7.62, 1.98, 5.10, 4.22, fill=MIST, line=LINE, radius=0.04)
    add_text(s, 7.86, 2.12, 4.5, 0.2, "EVIDENCE AT A GLANCE", size=7.8, color=TEAL, bold=True)
    stats = [("47 / 47", "automated tests passing"), ("6.21 s", "production build"),
             ("19", "interface captures from the running app")]
    for i, (num, label) in enumerate(stats):
        y = 2.44 + i * 0.92
        add_text(s, 7.86, y, 1.9, 0.42, num, size=18, color=INK, bold=True)
        add_text(s, 9.62, y + 0.06, 2.85, 0.34, label, size=8.8, color=BODY, line=1.15)
        if i < 2:
            rect(s, 7.86, y + 0.72, 4.6, 0.006, fill=LINE)
    chip(s, 7.86, 5.42, 2.55, 0.34, "Live demo · study-mapper.vercel.app",
         fill=TEAL, color=WHITE, size=7.6, bold=True)
    add_text(s, 7.86, 5.86, 4.4, 0.24, "Public repository · palturuvenubhaskar/StudyMapper",
             size=7.8, color=MUTED)

    rect(s, 0.62, 6.28, 12.10, 0.60, fill=TEAL_SOFT, line=None, radius=0.08)
    add_text(s, 0.84, 6.38, 11.7, 0.4,
             "Built at Anantha Lakshmi Institute of Technology & Sciences (Autonomous), Department of CSE, under the guidance of Ms. C. Sravani.",
             size=9.0, color=TEAL_DK, bold=True)
    notes(s, "Land the three pillars, one sentence each. Adaptive means the plan is an output of evidence, not a static document. Explainable means a student or examiner can trace why a topic is flagged weak or why it appears tomorrow — the data is in the same store. Local-first means the product survives a bad connection and treats cloud services as optional. The evidence panel repeats only measured facts. Close with the institution line, then move to future scope. Do not add new claims in the conclusion.")


def slide_16(prs):
    s = base(prs, "14 · Future scope", "Evidence first, then scale",
             "The next step is a classroom study; the roadmap follows from what that study reveals.",
             section="Close", page=16)
    horizons = [
        ("NOW → 3 MONTHS", TEAL, TEAL_SOFT, [
            ("Classroom pilot", "Pre/post quiz with a real cohort — the first test of learning impact."),
            ("Optional sync", "Account-based, conflict-safe sync without losing local speed."),
            ("Accessibility pass", "Keyboard paths, contrast and screen-reader labels."),
        ], "SIGNAL  ·  pre/post quiz gain in one cohort"),
        ("NEXT YEAR", BLUE, BLUE_SOFT, [
            ("Native mobile shell", "Capture small study sessions between classes."),
            ("Mentor view", "A cohort view for teachers — aggregate, not surveillance."),
            ("Content review queue", "A workflow for checking AI-generated material before sharing."),
        ], "SIGNAL  ·  weekly active students, review turnaround"),
        ("BEYOND", VIOLET, VIOLET_SOFT, [
            ("Multi-language content", "Localise notes and questions, not just the interface."),
            ("LMS integration", "LTI-style hand-off to institutional course spaces."),
            ("Longitudinal study", "A multi-semester evaluation with an education department."),
        ], "SIGNAL  ·  multi-semester retention and outcomes"),
    ]
    for i, (label, col, tint, items, signal) in enumerate(horizons):
        x = 0.62 + i * 4.10
        rect(s, x, 1.94, 3.92, 3.98, fill=WHITE, line=LINE, radius=0.04)
        rect(s, x, 1.94, 3.92, 0.50, fill=tint, radius=0.04)
        add_text(s, x + 0.20, 2.08, 3.4, 0.22, label, size=8.4, color=col, bold=True)
        for j, (t, d) in enumerate(items):
            y = 2.62 + j * 0.90
            add_text(s, x + 0.20, y, 3.5, 0.22, t, size=9.6, color=INK, bold=True)
            add_text(s, x + 0.20, y + 0.21, 3.5, 0.55, d, size=8.1, color=MUTED, line=1.14)
            if j < 2:
                rect(s, x + 0.20, y + 0.82, 3.5, 0.006, fill=LINE)
        rect(s, x + 0.20, 5.28, 3.5, 0.008, fill=LINE)
        add_text(s, x + 0.20, 5.36, 3.5, 0.18, "HOW WE WILL KNOW", size=6.6, color=col, bold=True)
        add_text(s, x + 0.20, 5.54, 3.5, 0.24, signal, size=7.6, color=BODY, bold=True)
    arrow(s, 0.62, 6.12, 12.20, 6.12, color=LINE, width=1.3)
    for i, (label, col, tint, items, signal) in enumerate(horizons):
        circle(s, 0.62 + i * 4.10 + 1.87, 6.03, 0.18, fill=col)
    add_text(s, 12.30, 6.02, 0.42, 0.2, "TIME", size=6.4, color=MUTED, bold=True)
    add_text(s, 0.62, 6.36, 12.0, 0.26,
             "Sequencing is deliberate: prove that the loop helps a student learn before adding surface area.",
             size=9.4, color=TEAL_DK, bold=True)
    notes(s, "The sequencing argument is the slide. The first horizon is not more features; it is a classroom pilot with a pre/post design, because that is the missing evidence. Optional sync and accessibility make the product usable beyond one browser. The next year is about reach and review: mobile, a mentor view that aggregates rather than surveils, and a workflow for checking AI content. Beyond that is institutional integration and a longitudinal study. If asked what is realistically next, answer the pilot and the accessibility pass — both are small enough to start immediately.")


def slide_17(prs):
    s = base(prs, "15 · References", "Sources used to verify this work",
             "Primary sources first; technical documentation only where it was needed to confirm a claim.",
             section="Close", page=17)
    refs_left = [
        ("1", "StudyMapper repository — github.com/palturuvenubhaskar/StudyMapper. Primary source for every implementation claim in this deck."),
        ("2", "Project abstract — “StudyMapper: AI-Driven Personalized Learning and Academic Management System”, ALITS, 2026–27."),
        ("3", "React 19 documentation — react.dev. Component model and hooks used across the 46 routed views."),
        ("4", "Vite documentation — vitejs.dev. Build tooling and the production bundle."),
        ("5", "Dexie.js documentation — dexie.org. IndexedDB schema, transactions and live queries."),
        ("6", "Firebase Authentication — firebase.google.com/docs/auth. Google, email and phone sign-in."),
    ]
    refs_right = [
        ("7", "OpenRouter API documentation — openrouter.ai/docs. Model routing, streaming and fallback behaviour."),
        ("8", "Pyodide documentation — pyodide.org. In-browser Python runtime used by coding practice."),
        ("9", "Tesseract.js — tesseract.js.org; pdf.js — mozilla.github.io/pdf.js; mammoth.js — mammothjs.org. File parsing pipeline."),
        ("10", "Mermaid — mermaid.js.org; KaTeX — katex.org. Diagram and mathematical rendering."),
        ("11", "P. A. Woźniak, “SuperMemo 2: Algorithm and implementation”, 1987. Basis of the SM-2 scheduler."),
        ("12", "Anki manual — apps.ankiweb.net. Reference implementation of spaced-repetition scheduling."),
    ]
    for px, pw, head in ((0.62, 5.86, "PRIMARY & PLATFORM"),
                         (6.86, 5.86, "LIBRARIES & METHODS")):
        rect(s, px, 1.90, pw, 4.62, fill=MIST, line=LINE, radius=0.035)
        add_text(s, px + 0.24, 2.04, pw - 0.48, 0.2, head, size=7.6, color=TEAL, bold=True)
    for col, entries, x in ((0, refs_left, 0.62), (1, refs_right, 6.86)):
        for i, (num, text) in enumerate(entries):
            y = 2.36 + i * 0.68
            add_text(s, x + 0.24, y - 0.01, 0.32, 0.22, num, size=8.8, color=TEAL, bold=True)
            add_text(s, x + 0.54, y - 0.01, 5.08, 0.64, text, size=8.4, color=BODY, line=1.16)
    add_text(s, 0.62, 6.64, 12.10, 0.24,
             "All web sources accessed October 2026. Only sources used to verify the implementation are listed.",
             size=8.2, color=MUTED, italic=True)
    notes(s, "Do not read the references. Explain the policy: the repository and the abstract are the primary sources; the rest are official documentation used to confirm specific technical behaviour. No secondary blog posts or unattributed claims are cited. If an examiner asks about the SM-2 formulation, point to Woźniak's original description and the Anki manual as the reference implementation. The repository link is also the fastest way for an evaluator to inspect the code after the presentation.")


def slide_18(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s, WHITE)
    rect(s, 0, 0, 0.10, 7.5, fill=TEAL)
    o = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(10.6), Inches(-1.5), Inches(4.4), Inches(4.4))
    o.fill.solid(); o.fill.fore_color.rgb = TEAL_SOFT; o.line.fill.background(); o.shadow.inherit = False
    add_text(s, 0.62, 1.75, 6.2, 0.85, "Thank you.", size=46, color=INK, bold=True)
    add_text(s, 0.62, 2.62, 6.2, 0.4, "Questions & discussion", size=18, color=TEAL, bold=True)
    rect(s, 0.62, 3.18, 0.9, 0.035, fill=TEAL)
    add_text(s, 0.62, 3.42, 6.0, 0.7,
             "The deck, the repository and the live build are available for review. Every screenshot in this presentation comes from the running application.",
             size=10.2, color=BODY, line=1.28)
    chip(s, 0.62, 4.36, 3.28, 0.38, "Live demo  ·  study-mapper.vercel.app", fill=TEAL, color=WHITE,
         size=8.8, bold=True)
    chip(s, 4.04, 4.36, 2.38, 0.38, "GitHub  ·  palturuvenubhaskar", fill=MIST, color=BODY, size=8.4,
         line=LINE, bold=True)
    add_text(s, 0.62, 6.52, 6.2, 0.3, "Academic project  ·  2026–27  ·  Batch 4",
             size=8.4, color=MUTED)

    rect(s, 7.32, 0.62, 5.40, 6.26, fill=MIST, line=LINE, radius=0.04)
    s.shapes.add_picture("assets/logo_clean.png", Inches(7.64), Inches(0.88), width=Inches(2.35))
    add_text(s, 7.64, 1.66, 4.6, 0.2, "TEAM", size=7.6, color=TEAL, bold=True)
    people = [
        ("P. Venubhaskar", "252G5A0520"), ("T. Vasanth", "252G5A0528"),
        ("V. Adarsh Kumar", "252G5A0532"), ("B. Sai Vikas", "252G5A0501"),
        ("C. Narendra", "252G5A0502"),
    ]
    for i, (name, roll) in enumerate(people):
        y = 1.88 + i * 0.44
        add_text(s, 7.64, y, 2.7, 0.24, name, size=9.6, color=INK, bold=True)
        add_text(s, 10.60, y + 0.02, 1.8, 0.22, roll, size=8.2, color=MUTED, align=PP_ALIGN.RIGHT)
        if i < 4:
            rect(s, 7.64, y + 0.36, 4.68, 0.005, fill=LINE)
    rect(s, 7.64, 4.10, 4.68, 0.008, fill=LINE)
    add_text(s, 7.64, 4.22, 4.6, 0.2, "GUIDE", size=7.6, color=TEAL, bold=True)
    add_text(s, 7.64, 4.42, 4.6, 0.26, "Ms. C. Sravani", size=11, color=INK, bold=True)
    add_text(s, 7.64, 4.78, 4.6, 0.2, "DEPARTMENT", size=7.6, color=TEAL, bold=True)
    add_text(s, 7.64, 4.98, 4.6, 0.5, "Computer Science & Engineering", size=10, color=INK, bold=True)
    add_text(s, 7.64, 5.40, 4.7, 0.2, "INSTITUTION", size=7.6, color=TEAL, bold=True)
    add_text(s, 7.64, 5.60, 4.7, 0.62,
             "Anantha Lakshmi Institute of Technology & Sciences (Autonomous)",
             size=9.6, color=INK, line=1.2)
    add_text(s, 7.64, 6.34, 4.7, 0.3, "Ananthapuramu, Andhra Pradesh", size=8.2, color=MUTED)
    notes(s, "Close by inviting questions rather than summarising again. Keep the two links visible: the live demo is the fastest way to show the loop, and the repository is where an examiner can verify implementation claims. The team panel includes roll numbers so the submission is identifiable. If time is short, offer a two-minute demo path: dashboard, planner, generated topic, flashcards. Then take questions, using the earlier slides as references for architecture and evidence.")
