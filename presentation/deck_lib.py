"""Design system + drawing helpers for the StudyMapper deck."""
import re
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from lxml import etree
from PIL import Image

FONT = "Calibri"

def C(h): return RGBColor.from_string(h)

INK      = C("152238")
BODY     = C("3D4A5C")
MUTED    = C("77859A")
TEAL     = C("0F7A82")
TEAL_DK  = C("0A565C")
TEAL_SOFT= C("E6F2F2")
MIST     = C("F5F8FA")
LINE     = C("DCE4EA")
WHITE    = C("FFFFFF")
AMBER    = C("B0781C")
AMBER_SOFT=C("FBF3E2")
ROSE     = C("A64B39")
ROSE_SOFT= C("FBEDE9")
GREEN    = C("2B7A58")
GREEN_SOFT=C("E9F4EE")
BLUE     = C("2E5C9A")
BLUE_SOFT= C("EAF0F8")
VIOLET   = C("5A4B87")
VIOLET_SOFT=C("EFECF7")

W, H = 13.333, 7.5


def set_bg(slide, color=WHITE):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _runs(p, text, size, color, bold=False, italic=False):
    for part in re.split(r"(\*\*.*?\*\*)", text):
        if not part:
            continue
        b = bold
        if part.startswith("**") and part.endswith("**"):
            part = part[2:-2]
            b = True
        r = p.add_run()
        r.text = part
        r.font.size = Pt(size)
        r.font.name = FONT
        r.font.color.rgb = color
        r.font.bold = b
        r.font.italic = italic


def add_text(slide, x, y, w, h, text, size=11.5, color=BODY, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line=1.12,
             space_after=0, italic=False, font=FONT, margins=(0, 0, 0, 0)):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    tf.margin_left, tf.margin_right, tf.margin_top, tf.margin_bottom = [Inches(m) for m in margins]
    lines = text.split("\n") if isinstance(text, str) else list(text)
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line
        p.space_after = Pt(space_after)
        _runs(p, ln, size, color, bold, italic)
    return box


def bullets(slide, x, y, w, h, items, size=11.5, color=BODY, accent=TEAL,
            gap=7, line=1.12, bold_lead=True):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = line
        p.space_after = Pt(gap)
        r = p.add_run()
        r.text = "▪  "
        r.font.size = Pt(size)
        r.font.name = FONT
        r.font.color.rgb = accent
        r.font.bold = True
        _runs(p, item, size, color, bold_lead)
    return box


def _strip_style(shp):
    for el in shp._element.findall(qn("p:style")):
        shp._element.remove(el)


def rect(slide, x, y, w, h, fill=WHITE, line=None, line_w=0.75,
         radius=None, dash=None):
    shp_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius is not None else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shp_type, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius is not None:
        try:
            s.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
        if dash:
            s.line.dash_style = dash
    s.shadow.inherit = False
    _strip_style(s)
    s.text_frame.word_wrap = True
    return s


def card(slide, x, y, w, h, fill=WHITE, line=LINE, radius=0.055, accent=None,
         accent_w=0.30):
    s = rect(slide, x, y, w, h, fill=fill, line=line, radius=radius)
    if accent:
        rect(slide, x + 0.018, y + 0.12, 0.045, h - 0.24, fill=accent)
    return s


def circle(slide, x, y, d, fill=TEAL, line=None, text=None, tcolor=WHITE,
           size=10, bold=True):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(0.75)
    s.shadow.inherit = False
    _strip_style(s)
    if text is not None:
        tf = s.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = FONT
        r.font.color.rgb = tcolor
    return s


def chip(slide, x, y, w, h, text, fill=MIST, color=BODY, size=9, bold=False,
         line=None, radius=0.5):
    s = rect(slide, x, y, w, h, fill=fill, line=line, radius=radius)
    tf = s.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.06)
    tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.name = FONT
    r.font.color.rgb = color
    r.font.bold = bold
    return s


def arrow(slide, x1, y1, x2, y2, color=MUTED, width=1.1, dash=None, head=True,
          head_size="med"):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                      Inches(x2), Inches(y2))
    conn.line.color.rgb = color
    conn.line.width = Pt(width)
    if dash:
        conn.line.dash_style = dash
    if head:
        ln = conn.line._get_or_add_ln()
        tail = etree.SubElement(ln, qn("a:tailEnd"))
        tail.set("type", "triangle")
        tail.set("w", head_size)
        tail.set("len", head_size)
    return conn


def chevron(slide, x, y, w, h, text, fill=TEAL, color=WHITE, size=10, first=False):
    shp = MSO_SHAPE.PENTAGON if first else MSO_SHAPE.CHEVRON
    s = slide.shapes.add_shape(shp, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    _strip_style(s)
    tf = s.text_frame
    tf.margin_left = Inches(0.18)
    tf.margin_right = Inches(0.10)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.name = FONT
    r.font.color.rgb = color
    return s


def cover_picture(slide, path, x, y, w, h, border=WHITE, border_w=1.0):
    iw, ih = Image.open(path).size
    src = iw / ih
    target = w / h
    pic = slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    if src > target:
        f = (1 - target / src) / 2
        pic.crop_left = f
        pic.crop_right = f
    else:
        f = (1 - src / target) / 2
        pic.crop_top = f
        pic.crop_bottom = f
    if border:
        pic.line.color.rgb = border
        pic.line.width = Pt(border_w)
    pic.shadow.inherit = False
    return pic


def framed_shot(slide, path, x, y, w, h, pad=0.07, caption=None, cap_size=8.5):
    rect(slide, x, y, w, h, fill=WHITE, line=LINE, radius=0.035)
    cover_picture(slide, path, x + pad, y + pad, w - 2 * pad, h - 2 * pad - (0.20 if caption else 0))
    if caption:
        add_text(slide, x + pad, y + h - 0.24, w - 2 * pad, 0.18, caption,
                 size=cap_size, color=MUTED, italic=True)


def base(prs, kicker, title, lead=None, section=None, page=None,
         title_size=27, bg=WHITE):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, bg)
    rect(slide, 0, 0, 0.10, 7.5, fill=TEAL)
    add_text(slide, 0.62, 0.36, 8.0, 0.24, kicker.upper(), size=9, color=TEAL, bold=True)
    add_text(slide, 0.62, 0.58, 11.2, 0.62, title, size=title_size, color=INK, bold=True, line=1.0)
    rect(slide, 0.62, 1.24, 0.85, 0.035, fill=TEAL)
    if lead:
        add_text(slide, 0.62, 1.36, 11.55, 0.50, lead, size=11.5, color=BODY, line=1.15)
    if section:
        add_text(slide, 9.2, 0.38, 3.5, 0.22, section.upper(), size=8, color=MUTED,
                 align=PP_ALIGN.RIGHT, bold=True)
    if page is not None:
        add_text(slide, 11.30, 7.12, 1.42, 0.2, f"{page:02d} / 18", size=8.4,
                 color=MUTED, align=PP_ALIGN.RIGHT)
    rect(slide, 0.62, 7.06, 12.10, 0.008, fill=LINE)
    add_text(slide, 0.62, 7.13, 8.0, 0.2,
             "StudyMapper  ·  AI-Driven Personalized Learning & Academic Management",
             size=8, color=MUTED)
    return slide


def notes(slide, text):
    tf = slide.notes_slide.notes_text_frame
    tf.text = text


# ---------- table helpers ----------
def _cell_border(cell, color="DCE4EA", width=0.75):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for el in tcPr.findall(qn(tag)):
            tcPr.remove(el)
    fills = [qn("a:solidFill"), qn("a:noFill"), qn("a:gradFill"), qn("a:pattFill")]
    idx = len(tcPr)
    for i, child in enumerate(tcPr):
        if child.tag in fills:
            idx = i
            break
    for tag in reversed(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
        ln = etree.Element(qn(tag))
        ln.set("w", str(int(width * 12700)))
        ln.set("cap", "flat"); ln.set("cmpd", "sng"); ln.set("algn", "ctr")
        sf = etree.SubElement(ln, qn("a:solidFill"))
        clr = etree.SubElement(sf, qn("a:srgbClr"))
        clr.set("val", color)
        etree.SubElement(ln, qn("a:prstDash")).set("val", "solid")
        tcPr.insert(idx, ln)


def table(slide, x, y, w, col_widths, rows, header=True, row_h=0.42,
          head_h=0.45, size=10, head_size=9.5, head_fill=TEAL,
          head_color=WHITE, body_color=BODY, zebra=MIST, bold_first=True):
    nrows, ncols = len(rows), len(rows[0])
    shp = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(w),
                                 Inches(head_h + row_h * (nrows - 1)))
    tbl = shp.table
    tbl.first_row = header
    tbl.horz_banding = False
    tbl.vert_banding = False
    # neutral style
    tblPr = tbl._tbl.tblPr
    for el in tblPr.findall(qn("a:tableStyleId")):
        tblPr.remove(el)
    sid = etree.SubElement(tblPr, qn("a:tableStyleId"))
    sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
    for i, wd in enumerate(col_widths):
        tbl.columns[i].width = Inches(wd)
    for r in range(nrows):
        tbl.rows[r].height = Inches(head_h if (r == 0 and header) else row_h)
        for c in range(ncols):
            cell = tbl.cell(r, c)
            cell.margin_left = Inches(0.09)
            cell.margin_right = Inches(0.09)
            cell.margin_top = Inches(0.04)
            cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if r == 0 and header:
                cell.fill.fore_color.rgb = head_fill
            elif zebra and r % 2 == 0:
                cell.fill.fore_color.rgb = zebra
            else:
                cell.fill.fore_color.rgb = WHITE
            _cell_border(cell, "DCE4EA", 0.6)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.line_spacing = 1.05
            txt = str(rows[r][c])
            bold = (r == 0 and header) or (bold_first and c == 0 and r > 0)
            _runs(p, txt, head_size if (r == 0 and header) else size,
                  head_color if (r == 0 and header) else body_color, bold)
    return tbl


def kv(slide, x, y, w, label, value, label_color=MUTED, size=10):
    add_text(slide, x, y, w, 0.2, label.upper(), size=7.8, color=label_color, bold=True)
    add_text(slide, x, y + 0.18, w, 0.45, value, size=size, color=INK, bold=True, line=1.0)
