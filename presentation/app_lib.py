"""App-coloured design system for the 13-slide StudyMapper deck.

The palette is taken from the running application: sky and indigo brand,
mint success, amber warning, rose danger, pastel aurora ground and white
glass cards.
"""
import os
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from lxml import etree
import deck_lib as D
from deck_lib import (add_text, bullets, rect, circle, chip, arrow,
                      cover_picture, notes, _strip_style)

# make every imported text helper use the app-like face
D.FONT = "Segoe UI"
FONT = "Segoe UI"

# ---------------------------------------------------------------- palette
INK        = RGBColor.from_string("0F172A")
BODY       = RGBColor.from_string("475569")
MUTED      = RGBColor.from_string("7C8AA3")
FAINT      = RGBColor.from_string("A9B5C9")
SKY        = RGBColor.from_string("0EA5E9")
SKY_DEEP   = RGBColor.from_string("0284C7")
SKY_SOFT   = RGBColor.from_string("E0F2FE")
INDIGO     = RGBColor.from_string("4F5BD5")
INDIGO_DK  = RGBColor.from_string("3730A3")
INDIGO_SOFT= RGBColor.from_string("EEF0FF")
VIOLET     = RGBColor.from_string("8B5CF6")
VIOLET_SOFT= RGBColor.from_string("F3EEFF")
MINT       = RGBColor.from_string("10B981")
MINT_DK    = RGBColor.from_string("047857")
MINT_SOFT  = RGBColor.from_string("E7FBF3")
AMBER      = RGBColor.from_string("F59E0B")
AMBER_DK   = RGBColor.from_string("B45309")
AMBER_SOFT = RGBColor.from_string("FFF7E6")
ROSE       = RGBColor.from_string("EF4444")
ROSE_DK    = RGBColor.from_string("B91C1C")
ROSE_SOFT  = RGBColor.from_string("FEECEC")
WHITE      = RGBColor.from_string("FFFFFF")
EDGE       = RGBColor.from_string("E3E8F5")
EDGE_2     = RGBColor.from_string("D5DCEE")
SHADOW     = RGBColor.from_string("E8ECF8")
CANVAS     = RGBColor.from_string("FBFCFE")
WASH       = RGBColor.from_string("F5F7FE")

W, H = 13.333, 7.5
MX = 0.70
CW = 11.93


# ---------------------------------------------------------------- basics
def bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = WHITE
    slide.shapes.add_picture("assets/bg_aurora.jpg", 0, 0,
                             width=Inches(W), height=Inches(H))


def txt(s, x, y, w, h, text, size=12, color=BODY, bold=False,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line=1.12,
        italic=False, space_after=0):
    return add_text(s, x, y, w, h, text, size=size, color=color, bold=bold,
                    align=align, anchor=anchor, line=line, italic=italic,
                    space_after=space_after)


def rich(s, x, y, w, h, parts, size=12, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, line=1.12):
    """parts = [(text, color, bold), ...]"""
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line
    for text, color, bold in parts:
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.name = FONT
        r.font.bold = bold
        r.font.color.rgb = color
    return box


def lab(s, x, y, w, text, color=INDIGO, size=9):
    return txt(s, x, y, w, 0.20, text.upper(), size=size, color=color,
               bold=True)


def card(s, x, y, w, h, fill=WHITE, line=EDGE, radius=0.06,
         shadow=True, accent=None, accent_w=0.05):
    if shadow:
        rect(s, x + 0.035, y + 0.05, w, h, fill=SHADOW, line=None, radius=radius)
    rect(s, x, y, w, h, fill=fill, line=line, radius=radius)
    if accent is not None:
        rect(s, x + 0.02, y + 0.14, accent_w, h - 0.28, fill=accent)


def grad(s, x, y, w, h, c1, c2, angle=0.0, radius=None):
    shape = MSO_SHAPE.ROUNDED_RECTANGLE if radius is not None else MSO_SHAPE.RECTANGLE
    shp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius is not None:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    shp.fill.gradient()
    shp.fill.gradient_angle = angle
    stops = shp.fill.gradient_stops
    stops[0].color.rgb = c1
    stops[0].position = 0.0
    stops[1].color.rgb = c2
    stops[1].position = 1.0
    shp.line.fill.background()
    shp.shadow.inherit = False
    _strip_style(shp)
    return shp


def icon(s, x, y, d, mono, color=INDIGO, fill=None, size=11):
    if fill is None:
        fill = WHITE
    rect(s, x, y, d, d, fill=fill, line=None, radius=0.28)
    txt(s, x, y + d / 2 - 0.12, d, 0.24, mono, size=size, color=color,
        bold=True, align=PP_ALIGN.CENTER)


def num_badge(s, x, y, d, text, color=INDIGO, fill=INDIGO_SOFT, size=10):
    circle(s, x, y, d, fill=fill, text=text, tcolor=color, size=size)


def pill(s, x, y, w, h, text, fill=INDIGO_SOFT, color=INDIGO, size=9.5,
         line=None, bold=True):
    return chip(s, x, y, w, h, text, fill=fill, color=color, size=size,
                bold=bold, line=line, radius=0.5)


def shot_card(s, path, x, y, w, img_h, title, sub, accent=INDIGO):
    h = img_h + 0.58
    card(s, x, y, w, h, fill=WHITE, line=EDGE, radius=0.05)
    cover_picture(s, path, x + 0.07, y + 0.07, w - 0.14, img_h, border=None)
    rect(s, x + 0.07, y + 0.07 + img_h + 0.07, w - 0.14, 0.028, fill=accent)
    txt(s, x + 0.10, y + img_h + 0.14, w - 0.18, 0.22, title, size=10.5,
        color=INK, bold=True)
    txt(s, x + 0.10, y + img_h + 0.36, w - 0.18, 0.20, sub, size=8.6,
        color=MUTED)
    return h


def arr(s, x1, y1, x2, y2, color=MUTED, width=1.1, dash=None, head=True):
    c = arrow(s, x1, y1, x2, y2, color=color, width=width, dash=dash, head=head)
    c.shadow.inherit = False
    _strip_style(c)
    return c


def stick(s, cx, top, scale=1.0, color=INK):
    d = 0.20 * scale
    circle(s, cx - d / 2, top, d, fill=color)
    arr(s, cx, top + d, cx, top + 0.46 * scale, color=color, width=1.5, head=False)
    arr(s, cx - 0.19 * scale, top + 0.20 * scale, cx + 0.19 * scale,
        top + 0.20 * scale, color=color, width=1.5, head=False)
    arr(s, cx, top + 0.46 * scale, cx - 0.15 * scale, top + 0.70 * scale,
        color=color, width=1.5, head=False)
    arr(s, cx, top + 0.46 * scale, cx + 0.15 * scale, top + 0.70 * scale,
        color=color, width=1.5, head=False)


def diamond(s, cx, cy, w, h, fill=AMBER_SOFT, line=AMBER, text="", size=9,
            color=INK):
    shp = s.shapes.add_shape(MSO_SHAPE.DIAMOND, Inches(cx - w / 2),
                             Inches(cy - h / 2), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = line
    shp.line.width = Pt(1.1)
    shp.shadow.inherit = False
    _strip_style(shp)
    if text:
        tf = shp.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = True
        r.font.name = FONT
        r.font.color.rgb = color
    return shp


def elbow(s, pts, color=INDIGO, width=1.4, dash=None):
    for i in range(len(pts) - 1):
        x1, y1 = pts[i]
        x2, y2 = pts[i + 1]
        arr(s, x1, y1, x2, y2, color=color, width=width, dash=dash,
            head=(i == len(pts) - 2))


# ---------------------------------------------------------------- page frame
def base(prs, num, kicker, title, lead=None, title_size=33):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    grad(s, MX, 0.45, 0.18, 0.18, SKY, INDIGO, angle=35, radius=0.25)
    txt(s, MX + 0.27, 0.40, 2.0, 0.24, "StudyMapper", size=9.5, color=INK, bold=True)
    txt(s, 6.6, 0.41, 6.03, 0.22, kicker.upper(), size=9.5, color=INDIGO,
        bold=True, align=PP_ALIGN.RIGHT)
    txt(s, MX, 0.72, CW, 0.62, title, size=title_size, color=INK, bold=True, line=0.98)
    if lead:
        txt(s, MX, 1.36, 11.45, 0.48, lead, size=14.5, color=BODY, line=1.14)
    rect(s, MX, 7.04, CW, 0.008, fill=EDGE_2)
    txt(s, MX, 7.13, 8.6, 0.18,
        "StudyMapper  ·  AI-Driven Personalized Learning & Academic Management",
        size=8, color=MUTED)
    txt(s, 11.0, 7.12, 1.63, 0.20, f"{num:02d} / 13", size=8.8, color=MUTED,
        bold=True, align=PP_ALIGN.RIGHT)
    return s
