"""Slides 1-6."""
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from deck_lib import *


def slide_01(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s, WHITE)
    rect(s, 0, 0, 0.10, 7.5, fill=TEAL)
    # soft background shapes
    o = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(10.4), Inches(-1.7), Inches(4.8), Inches(4.8))
    o.fill.solid(); o.fill.fore_color.rgb = TEAL_SOFT; o.line.fill.background(); o.shadow.inherit = False
    o2 = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(12.2), Inches(5.9), Inches(1.9), Inches(1.9))
    o2.fill.solid(); o2.fill.fore_color.rgb = AMBER_SOFT; o2.line.fill.background(); o2.shadow.inherit = False

    s.shapes.add_picture("assets/logo_clean.png", Inches(0.62), Inches(0.48), width=Inches(2.55))

    add_text(s, 0.62, 1.62, 6.4, 0.24, "ACADEMIC PROJECT PRESENTATION  ·  2026–27", size=9, color=TEAL, bold=True)
    add_text(s, 0.62, 1.86, 6.5, 0.85, "StudyMapper", size=46, color=INK, bold=True, line=0.95)
    add_text(s, 0.62, 2.72, 6.15, 0.95,
             "AI-driven personalised learning and academic management — plan, learn, practise and improve in one workspace.",
             size=15, color=BODY, line=1.18)
    rect(s, 0.62, 3.78, 0.9, 0.035, fill=TEAL)

    add_text(s, 0.62, 3.98, 6.3, 0.24, "TEAM", size=7.8, color=MUTED, bold=True)
    add_text(s, 0.62, 4.18, 6.4, 0.3,
             "P. Venubhaskar  ·  T. Vasanth  ·  V. Adarsh Kumar  ·  B. Sai Vikas  ·  C. Narendra",
             size=10.5, color=INK, bold=True)
    add_text(s, 0.62, 4.56, 6.3, 0.24, "GUIDE  ·  Ms. C. Sravani", size=10.5, color=TEAL, bold=True)
    add_text(s, 0.62, 4.90, 6.3, 0.5,
             "Department of Computer Science & Engineering\nAnantha Lakshmi Institute of Technology & Sciences (Autonomous)",
             size=9.5, color=MUTED, line=1.25)

    chip(s, 0.62, 5.72, 3.30, 0.36, "Live demo  ·  study-mapper.vercel.app",
         fill=TEAL, color=WHITE, size=9, bold=True)
    chip(s, 4.06, 5.72, 1.75, 0.36, "47 / 47 tests", fill=GREEN_SOFT, color=GREEN, size=9, bold=True)
    add_text(s, 0.62, 6.28, 6.2, 0.5,
             "Everything shown in this deck is drawn from the project repository and the running application.",
             size=8.5, color=MUTED, italic=True, line=1.2)

    # right panel
    rect(s, 7.38, 0.62, 5.34, 6.28, fill=MIST, line=LINE, radius=0.035)
    framed_shot(s, "assets/dashboard-light.png", 7.66, 0.90, 4.78, 2.98,
                caption="Dashboard — the daily workspace")
    framed_shot(s, "assets/planner-schedule.png", 7.66, 4.28, 2.28, 1.52, pad=0.05,
                caption="Adaptive planner")
    framed_shot(s, "assets/topic-diagram.png", 10.16, 4.28, 2.28, 1.52, pad=0.05,
                caption="Generated diagram")
    add_text(s, 7.66, 6.08, 4.78, 0.5,
             "Genuine captures from the application at 1440 × 900, light theme. No mockups.",
             size=8, color=MUTED, italic=True, line=1.2)

    notes(s, "Open by naming the project and the problem space, not the technology. StudyMapper is a browser-based study workspace for a single student: it turns a syllabus into an adaptive plan, generates study material, and measures what needs revision. The three captures on the right are the actual application — dashboard, planner and a generated Mermaid diagram. Mention the team and guide, then say the deck moves from why the system exists, to how it works, to what has actually been verified. Keep this to about forty seconds; the next slide explains the structure.")


def slide_02(prs):
    s = base(prs, "Agenda", "Eighteen slides, five acts",
             "Why the system exists, what it does, how it is built, what is verified, and where it goes next.",
             page=2)
    acts = [
        ("01", "Foundations", "03–06", "Abstract, problem, objectives and existing work.",
         ["Context & motivation", "Problem definition", "Objectives & scope", "Literature review"],
         "You will know the problem, the scope and the gap."),
        ("02", "The system", "07–09", "The product loop, its architecture and the method.",
         ["Proposed modules", "Layered architecture", "Development method"],
         "You will see how the nine modules connect."),
        ("03", "Engineering", "10–12", "The decisions that make it work in a browser.",
         ["Technology choices", "Data model & flow", "Interface design"],
         "You will see the data model and interface decisions."),
        ("04", "Evidence", "13–14", "What was tested, and what is not claimed.",
         ["Verified results", "Contributions", "Limitations"],
         "You will know exactly what was tested — and what was not."),
        ("05", "Close", "15–18", "Conclusion, next steps and sources.",
         ["Conclusion", "Future scope", "References", "Questions"],
         "You will know the next steps and the sources."),
    ]
    fills = [TEAL, TEAL_DK, TEAL, TEAL_DK, TEAL]
    for i, (num, name, rng, desc, items, outcome) in enumerate(acts):
        x = 0.62 + i * 2.38
        chevron(s, x, 1.92, 2.55, 0.48, f"{num}  {name.upper()}", fill=fills[i],
                size=8.5, first=(i == 0))
        card(s, x, 2.52, 2.28, 3.18, fill=WHITE, line=LINE, accent=fills[i])
        add_text(s, x + 0.20, 2.74, 1.9, 0.2, f"SLIDES {rng}", size=7.8, color=fills[i], bold=True)
        add_text(s, x + 0.20, 2.98, 1.95, 0.55, desc, size=9.4, color=MUTED, line=1.18)
        add_text(s, x + 0.20, 3.64, 1.95, 1.0, "\n".join("·  " + t for t in items),
                 size=8.8, color=BODY, line=1.5)
        rect(s, x + 0.20, 4.62, 1.88, 0.006, fill=LINE)
        add_text(s, x + 0.20, 4.70, 1.9, 0.18, "OUTCOME", size=6.6, color=fills[i], bold=True)
        add_text(s, x + 0.20, 4.88, 1.92, 0.7, outcome, size=8.4, color=BODY, line=1.18)

    rect(s, 0.62, 5.92, 12.10, 0.74, fill=MIST, line=LINE, radius=0.06)
    add_text(s, 0.86, 6.02, 2.2, 0.2, "HOW TO READ THIS DECK", size=8, color=TEAL, bold=True)
    add_text(s, 0.86, 6.21, 11.6, 0.34,
             "Teal marks the system, amber marks evidence, rose marks a limitation. Implementation claims come from the repository; every screenshot is from the running app.",
             size=9.5, color=BODY, line=1.2)
    notes(s, "Use this slide as a map, not a list. The five acts answer five questions in order: why does this need to exist, what does it do, how is it engineered, what can we actually prove, and what comes next. Point out the colour convention in the bottom band — it is used consistently, so the audience can tell a design decision from a verified result from a limitation. Say that the deep technical detail is concentrated in acts three and four, which lets non-technical listeners follow the narrative without losing the engineers.")


def slide_03(prs):
    s = base(prs, "01 · Abstract", "A study loop that responds to the learner",
             "StudyMapper turns a syllabus into an adaptive plan, generates study material, and measures what actually needs revision — in one browser-based workspace.",
             section="Foundations", page=3)
    add_text(s, 0.62, 1.98, 5.9, 0.2, "IN ONE PARAGRAPH", size=8, color=TEAL, bold=True)
    add_text(s, 0.62, 2.20, 5.90, 1.15,
             "The system combines curriculum management, AI-assisted content generation, spaced-repetition scheduling and practice analytics. A student builds or imports a syllabus, receives a difficulty-aware study plan, studies generated notes and diagrams, practises with question banks, and sees weak topics recomputed from recent activity. Data lives locally in the browser; AI and sign-in are optional services.",
             size=10.4, color=BODY, line=1.25)

    add_text(s, 0.62, 3.58, 5.9, 0.2, "WHY IT MATTERED", size=8, color=TEAL, bold=True)
    rows = [
        ("Study time is fixed; attention is not.",
         "A plan that ignores difficulty and weak topics wastes the hours a student does have."),
        ("Feedback arrives after the exam.",
         "Continuous practice data makes weakness visible while it can still be acted on."),
        ("Study tools are scattered.",
         "Notes, plans, tests and progress sit in different apps, so context is lost between them."),
    ]
    for i, (t, d) in enumerate(rows):
        y = 3.84 + i * 0.78
        rect(s, 0.62, y + 0.04, 0.09, 0.09, fill=TEAL)
        add_text(s, 0.86, y - 0.02, 5.65, 0.24, t, size=10.8, color=INK, bold=True)
        add_text(s, 0.86, y + 0.20, 5.65, 0.45, d, size=9.2, color=MUTED, line=1.16)

    add_text(s, 0.62, 6.18, 5.9, 0.22,
             "DESIGN PRINCIPLE  ·  The plan is an output of evidence, not a document.",
             size=8.2, color=TEAL, bold=True)
    add_text(s, 0.62, 6.50, 6.0, 0.4,
             "Built as a single-page React application. No server is required for the core study loop.",
             size=8.5, color=MUTED, italic=True)

    # loop diagram
    rect(s, 6.92, 1.95, 5.78, 4.67, fill=MIST, line=LINE, radius=0.03)
    add_text(s, 7.18, 2.12, 5.3, 0.2, "THE LEARNING LOOP", size=8, color=TEAL, bold=True)
    nodes = [
        (8.88, 2.38, "PLAN", "Exam date, difficulty, weak scores", TEAL),
        (10.61, 3.84, "LEARN", "Notes, diagrams, imported PDFs", BLUE),
        (8.88, 5.30, "PRACTISE", "Question banks, mock exams, code", AMBER),
        (7.14, 3.84, "IMPROVE", "Weakness score, SM-2 review, re-plan", GREEN),
    ]
    for (nx, ny, name, desc, col) in nodes:
        card(s, nx, ny, 2.02, 0.92, fill=WHITE, line=LINE, accent=col)
        add_text(s, nx + 0.18, ny + 0.20, 1.75, 0.22, name, size=9.6, color=INK, bold=True)
        add_text(s, nx + 0.18, ny + 0.44, 1.72, 0.44, desc, size=7.5, color=MUTED, line=1.1)
    circle(s, 9.40, 3.62, 0.95, fill=WHITE, line=LINE)
    add_text(s, 9.42, 3.84, 0.91, 0.5, "ONE\nLOOP", size=8.6, color=TEAL, bold=True,
             align=PP_ALIGN.CENTER, line=1.0)
    arrow(s, 10.16, 3.33, 10.80, 3.80, color=TEAL, width=1.2)
    arrow(s, 11.62, 4.76, 10.98, 5.26, color=TEAL, width=1.2)
    arrow(s, 8.88, 5.76, 8.24, 5.12, color=TEAL, width=1.2)
    arrow(s, 8.16, 3.84, 8.80, 3.34, color=TEAL, width=1.2)
    add_text(s, 7.18, 6.38, 5.3, 0.22,
             "Each pass writes back to the same store, so the next plan starts from evidence.",
             size=8.0, color=MUTED, italic=True)
    notes(s, "Read the abstract slowly enough that the non-technical listeners catch the loop, then use the diagram to make it concrete. Plan takes the exam date, the difficulty of each topic and the current weakness scores; Learn is where generated notes and diagrams live; Practise is question banks, mock exams and in-browser coding; Improve turns attempts into weakness scores, schedules flashcards with SM-2 and feeds the next plan. Stress that this is one data store — the arrow returning to Plan is not a metaphor, it is the same records being recomputed.")


def slide_04(prs):
    s = base(prs, "02 · Problem", "The gap is not content — it is response",
             "Students rarely fail for lack of material. They fail because nothing in their day responds to what they actually understand.",
             section="Foundations", page=4)
    cards = [
        ("01", "Fragmented workflow",
         "Notes live in one app, deadlines in another, tests in a third. Context is lost every time a student switches.",
         "A student rewrites the same notes in three tools.",
         "Revision restarts from zero."),
        ("02", "Static study plans",
         "A plan made once ignores topic difficulty, weak areas and slipping dates. It stays comfortable while the exam gets closer.",
         "The weakest topic gets the same 30 minutes as the strongest.",
         "Time goes where it is easiest, not where it is needed."),
        ("03", "Feedback after the fact",
         "Progress becomes visible only when marks arrive. Misconceptions survive weeks of practice unchallenged.",
         "A wrong method is repeated until the mock exam.",
         "The same mistakes repeat."),
    ]
    for i, (num, title, body, example, cost) in enumerate(cards):
        x = 0.62 + i * 4.10
        card(s, x, 2.02, 3.95, 3.28, fill=WHITE, line=LINE, accent=ROSE)
        add_text(s, x + 0.24, 2.24, 0.6, 0.25, num, size=9, color=ROSE, bold=True)
        add_text(s, x + 0.24, 2.50, 3.45, 0.5, title, size=14, color=INK, bold=True)
        add_text(s, x + 0.24, 3.02, 3.48, 0.6, body, size=9.8, color=BODY, line=1.24)
        add_text(s, x + 0.24, 3.66, 3.45, 0.5, example, size=8.7, color=MUTED, italic=True, line=1.18)
        rect(s, x + 0.24, 4.35, 3.42, 0.008, fill=LINE)
        add_text(s, x + 0.24, 4.44, 3.4, 0.2, "COST TO THE STUDENT", size=7.2, color=MUTED, bold=True)
        add_text(s, x + 0.24, 4.62, 3.4, 0.5, cost, size=9.4, color=ROSE, bold=True, line=1.15)

    rect(s, 0.62, 5.55, 12.10, 1.12, fill=MIST, line=LINE, radius=0.06)
    add_text(s, 0.86, 5.72, 2.4, 0.2, "DESIGN IMPLICATION", size=8, color=TEAL, bold=True)
    add_text(s, 0.86, 5.94, 2.35, 0.5,
             "Three failures define\nthe requirements.", size=9.5, color=BODY, line=1.2)
    steps = [("One workspace", "context survives"), ("A plan that recomputes", "from real evidence"),
             ("Measurement inside practice", "not after it")]
    for i, (t, d) in enumerate(steps):
        x = 3.30 + i * 3.15
        if i > 0:
            arrow(s, x - 0.30, 6.11, x - 0.04, 6.11, color=TEAL, width=1.3)
        add_text(s, x, 5.79, 2.75, 0.24, t, size=10.2, color=INK, bold=True)
        add_text(s, x, 6.04, 2.75, 0.22, d, size=8.6, color=MUTED)
    notes(s, "Frame the problem in student terms before any technology. The first card is about context switching: notes, deadlines and tests are separate, so every switch costs momentum. The second is about plans that never react — a student can follow the schedule perfectly and still under-prepare the weak topic. The third is delayed feedback: by the time marks arrive, the misconception has been practised for weeks. The bottom band converts these into three requirements that the next slides answer. Avoid blaming students; the point is that the tools do not respond.")


def slide_05(prs):
    s = base(prs, "03 · Objectives", "Four objectives, each tied to something observable",
             "Scope was deliberately narrow: the individual study loop for a CSE undergraduate — not institutional administration.",
             section="Foundations", page=5)
    rows = [
        ("01", "Make planning adaptive",
         "Regenerate the schedule from exam date, topic difficulty and weak scores.",
         "A plan that visibly changes after practice."),
        ("02", "Turn source material into study assets",
         "Generate notes, diagrams, questions and flashcards from a syllabus or an imported PDF.",
         "Artefacts created and stored inside the app."),
        ("03", "Measure understanding continuously",
         "Score every attempt and expose weak topics rather than a single average.",
         "A 0–100 weakness score per topic."),
        ("04", "Make consistency sustainable",
         "Use streaks, XP, quests and small daily goals to reduce drop-off.",
         "An active streak and completed daily quests."),
    ]
    for i, (num, t, d, m) in enumerate(rows):
        y = 1.98 + i * 1.05
        circle(s, 0.64, y + 0.06, 0.42, fill=TEAL_SOFT, text=num, tcolor=TEAL, size=9)
        add_text(s, 1.24, y, 7.2, 0.28, t, size=13, color=INK, bold=True)
        add_text(s, 1.24, y + 0.30, 7.25, 0.45, d, size=10.2, color=BODY, line=1.18)
        add_text(s, 1.24, y + 0.72, 7.2, 0.2, "Measured by  ·  " + m, size=8.4, color=TEAL, bold=True)
        if i < 3:
            rect(s, 1.24, y + 0.98, 7.2, 0.008, fill=LINE)
    add_text(s, 0.62, 6.44, 7.8, 0.3,
             "Each objective has an observable check — that is exactly what the evidence slide reports.",
             size=8.8, color=MUTED, italic=True)

    rect(s, 8.80, 1.96, 3.92, 4.72, fill=MIST, line=LINE, radius=0.03)
    add_text(s, 9.02, 2.14, 3.4, 0.2, "SCOPE", size=8, color=TEAL, bold=True)
    add_text(s, 9.02, 2.38, 3.4, 0.22, "IN SCOPE", size=8, color=GREEN, bold=True)
    add_text(s, 9.02, 2.60, 3.45, 1.35,
             "·  CSE study loop: DBMS, OS, DSA\n·  Browser, offline-first\n·  Individual student\n·  AI optional, not required",
             size=9.4, color=BODY, line=1.32)
    rect(s, 9.02, 4.02, 3.42, 0.008, fill=LINE)
    add_text(s, 9.02, 4.14, 3.4, 0.22, "OUT OF SCOPE", size=8, color=ROSE, bold=True)
    add_text(s, 9.02, 4.36, 3.45, 1.0,
             "·  Institutional timetabling\n·  Grading or proctored exams\n·  Native mobile application",
             size=9.4, color=BODY, line=1.32)
    add_text(s, 9.02, 5.44, 3.42, 0.2, "SUCCESS CRITERIA", size=8, color=TEAL, bold=True)
    add_text(s, 9.02, 5.66, 3.5, 0.85,
             "47/47 automated tests  ·  production build  ·  live deployment  ·  demo dataset",
             size=9, color=INK, bold=True, line=1.22)
    notes(s, "Each objective has a measurement, which is the point of the slide. Adaptive planning is not “we added AI”; it is that the schedule changes when the evidence changes. Content generation is measured by artefacts that exist in the app, not by the model responding. Continuous measurement is a per-topic score, not one average. Consistency is streaks and quests, which the interface treats as a first-class feature. The right panel is equally important: saying what is out of scope protects the project from over-claiming. Institutional timetabling, formal grading and a native app were deliberately excluded.")


def slide_06(prs):
    s = base(prs, "04 · Existing work", "Five strands inform the design — none closes the loop",
             "The review focused on what each strand does well and what it leaves unresolved for an individual learner.",
             section="Foundations", page=6)
    rows = [
        ("Strand", "What it does well", "What it leaves unresolved"),
        ("Adaptive learning platforms", "Sequence content and mastery paths; strong structured curricula",
         "Personal scheduling is weak; the plan rarely reacts to a single student's week"),
        ("AI tutors & content generation", "Explain, summarise and generate questions on demand",
         "Usually cloud-only and disconnected from the learner's own stored data"),
        ("LMS analytics", "Cohort and course dashboards; institutional reporting",
         "An institutional view, not a decision about the next hour of study"),
        ("Spaced repetition", "SM-2 schedules review at the edge of recall",
         "Isolated from the syllabus, exam dates and the rest of the workflow"),
        ("Gamified learning", "XP, streaks and quests sustain habit and return visits",
         "Engagement without adaptation; points do not change what is studied"),
    ]
    table(s, 0.62, 1.94, 12.10, [2.45, 4.55, 5.10], rows, row_h=0.64, head_h=0.44, size=9.4)
    rect(s, 0.62, 5.92, 12.10, 0.80, fill=MIST, line=LINE, radius=0.06)
    add_text(s, 0.86, 6.04, 1.3, 0.2, "POSITIONING", size=7.8, color=TEAL, bold=True)
    add_text(s, 0.86, 6.24, 11.5, 0.34,
             "StudyMapper is not another LMS. It is a personal layer that connects planning, content, practice and measurement — and keeps working offline.",
             size=9.8, color=BODY)
    add_text(s, 7.2, 6.04, 5.3, 0.2,
             "LOCAL-FIRST   ·   AI OPTIONAL   ·   CLOSED LOOP",
             size=7.2, color=TEAL, bold=True, align=PP_ALIGN.RIGHT)
    notes(s, "Do not read the table row by row; pick the contrast. Adaptive platforms are strong on content but weak on a personal weekly plan. AI tutors generate well but usually cannot see the learner's stored data. LMS analytics serve the institution, not the next hour. Spaced repetition is excellent but isolated. Gamification sustains habit without changing what is studied. StudyMapper's claim is modest and precise: it connects these strands in one local store. The three tags at the bottom summarise the positioning — local-first, AI optional, closed feedback loop.")
